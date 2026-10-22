# /// script
# requires-python = ">=3.10"
# dependencies = ["openpyxl==3.1.5"]
# ///
import argparse
import csv
import datetime
import io
import json
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries


def render(value: object) -> object:
    if value is None:
        return ""
    if isinstance(value, datetime.datetime):
        return value.isoformat(sep=" ", timespec="seconds").replace(" 00:00:00", "")
    if isinstance(value, (datetime.date, datetime.time)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer() and abs(value) < 1e15:
        return int(value)
    return value


def read_csv(path: str, delimiter: str | None) -> list[list[str]]:
    text = Path(path).read_text(encoding="utf-8-sig", errors="replace")
    dialect_delimiter = delimiter
    if dialect_delimiter is None:
        try:
            dialect_delimiter = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|").delimiter
        except csv.Error:
            dialect_delimiter = ","
    return [row for row in csv.reader(io.StringIO(text), delimiter=dialect_delimiter)]


def main() -> int:
    parser = argparse.ArgumentParser(description="Read cells from a workbook or CSV file.")
    parser.add_argument("file")
    parser.add_argument("--sheet", help="sheet name; the active sheet when omitted")
    parser.add_argument("--range", dest="cell_range", help="e.g. A1:H40")
    parser.add_argument("--start-row", type=int, default=None, help="1-based first row of the window")
    parser.add_argument("--limit", type=int, default=200, help="rows per window")
    parser.add_argument("--format", choices=["markdown", "csv", "json"], default="markdown")
    parser.add_argument("--formulas", action="store_true", help="show formulas instead of cached values")
    parser.add_argument("--delimiter", default=None, help="CSV input delimiter")
    parser.add_argument("--output", help="write to this file instead of stdout")
    args = parser.parse_args()

    if args.file.lower().endswith((".csv", ".tsv", ".txt")):
        rows = read_csv(args.file, "\t" if args.file.lower().endswith(".tsv") else args.delimiter)
        first_row, first_col = 1, 1
        total_rows = len(rows)
        title = Path(args.file).name
    else:
        book = load_workbook(args.file, data_only=not args.formulas, read_only=False)
        sheet = book[args.sheet] if args.sheet else book.active
        title = sheet.title
        if args.cell_range:
            min_col, min_row, max_col, max_row = range_boundaries(args.cell_range)
        else:
            min_col, min_row, max_col, max_row = 1, 1, sheet.max_column, sheet.max_row
        total_rows = max_row
        start = max(args.start_row or min_row, min_row)
        stop = min(max_row, start + max(args.limit, 1) - 1) if not args.output else max_row
        rows = [
            [render(cell.value) for cell in row]
            for row in sheet.iter_rows(min_row=start, max_row=stop, min_col=min_col, max_col=max_col)
        ]
        first_row, first_col = start, min_col

    if args.file.lower().endswith((".csv", ".tsv", ".txt")) and not args.output:
        start = max(args.start_row or 1, 1)
        rows = rows[start - 1 : start - 1 + max(args.limit, 1)]
        first_row = start

    while rows and all(value == "" for value in rows[-1]):
        rows.pop()
    width = max((len(row) for row in rows), default=0)
    while width > 0 and all(len(row) < width or row[width - 1] == "" for row in rows):
        width -= 1
    rows = [row[:width] + [""] * (width - len(row[:width])) for row in rows]

    if args.format == "csv":
        buffer = io.StringIO()
        csv.writer(buffer).writerows(rows)
        text = buffer.getvalue()
    elif args.format == "json":
        text = json.dumps(
            {"sheet": title, "firstRow": first_row, "firstColumn": get_column_letter(first_col), "rows": rows},
            ensure_ascii=False,
            default=str,
        )
    else:
        header = [""] + [get_column_letter(first_col + index) for index in range(width)]
        lines = ["| " + " | ".join(header) + " |", "|" + " --- |" * len(header)]
        for offset, row in enumerate(rows):
            cells = [str(value).replace("|", "\\|").replace("\n", " ") for value in row]
            lines.append(f"| {first_row + offset} | " + " | ".join(cells) + " |")
        text = f"Sheet: {title}\n\n" + "\n".join(lines) + "\n"

    if args.output:
        encoding = "utf-8-sig" if args.format == "csv" else "utf-8"
        Path(args.output).write_text(text, encoding=encoding, newline="")
        print(f"Wrote {len(rows)} rows -> {Path(args.output).resolve()}")
        return 0
    sys.stdout.write(text)
    last = first_row + len(rows) - 1
    if rows and last < total_rows:
        print(f"\n[rows {first_row}-{last} of {total_rows}; continue with --start-row {last + 1}]")
    return 0


if __name__ == "__main__":
    sys.exit(main())
