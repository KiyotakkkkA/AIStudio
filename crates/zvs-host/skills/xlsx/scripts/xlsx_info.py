# /// script
# requires-python = ">=3.10"
# dependencies = ["openpyxl==3.1.5"]
# ///
import argparse
import json
import sys

from openpyxl import load_workbook


def main() -> int:
    parser = argparse.ArgumentParser(description="Print the structure of a workbook as JSON.")
    parser.add_argument("file")
    args = parser.parse_args()

    formulas_book = load_workbook(args.file, data_only=False)
    values_book = load_workbook(args.file, data_only=True)
    sheets = []
    total_formulas = 0
    uncalculated = 0
    for sheet in formulas_book.worksheets:
        values = values_book[sheet.title]
        formulas = 0
        empty_results = 0
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    formulas += 1
                    if values[cell.coordinate].value is None:
                        empty_results += 1
        total_formulas += formulas
        uncalculated += empty_results
        sheets.append(
            {
                "name": sheet.title,
                "state": sheet.sheet_state,
                "usedRange": sheet.dimensions,
                "rows": sheet.max_row,
                "columns": sheet.max_column,
                "formulas": formulas,
                "mergedCells": [str(item) for item in sheet.merged_cells.ranges][:30],
                "freeze": sheet.freeze_panes,
                "autoFilter": sheet.auto_filter.ref,
                "tables": {name: table.ref for name, table in sheet.tables.items()},
                "charts": len(getattr(sheet, "_charts", [])),
                "images": len(getattr(sheet, "_images", [])),
                "dataValidations": len(sheet.data_validations.dataValidation),
                "conditionalFormats": len(sheet.conditional_formatting),
            }
        )
    names = {}
    for name, definition in formulas_book.defined_names.items():
        names[name] = definition.attr_text
    info = {
        "sheets": sheets,
        "activeSheet": formulas_book.active.title if formulas_book.active else None,
        "definedNames": names,
        "formulas": total_formulas,
        "calculated": total_formulas == 0 or uncalculated < total_formulas,
        "formulasWithoutCachedValue": uncalculated,
    }
    print(json.dumps(info, ensure_ascii=False, indent=2, default=str))
    return 0


if __name__ == "__main__":
    sys.exit(main())
