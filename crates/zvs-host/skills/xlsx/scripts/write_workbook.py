# /// script
# requires-python = ">=3.10"
# dependencies = ["openpyxl==3.1.5"]
# ///
import argparse
import csv
import datetime
from copy import copy
import io
import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, PatternFill, Side
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.cell import coordinate_from_string, range_boundaries
from openpyxl.workbook.defined_name import DefinedName

ISO_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
ISO_DATETIME = re.compile(r"^\d{4}-\d{2}-\d{2}[ T]\d{2}:\d{2}(:\d{2})?$")
NUMBER = re.compile(r"^-?\d+([.,]\d+)?$")
THIN = Side(style="thin", color="A6A6A6")
HEADER_FILL = PatternFill("solid", fgColor="E7E6E6")


def convert(value: object, parse_numbers: bool) -> object:
    if not isinstance(value, str):
        return value
    if ISO_DATE.match(value):
        return datetime.date.fromisoformat(value)
    if ISO_DATETIME.match(value):
        return datetime.datetime.fromisoformat(value.replace(" ", "T"))
    if parse_numbers and NUMBER.match(value.strip()):
        text = value.strip().replace(",", ".")
        return float(text) if "." in text else int(text)
    return value


def put(sheet, coordinate: str, value: object, parse_numbers: bool = False) -> None:
    cell = sheet[coordinate]
    converted = convert(value, parse_numbers)
    cell.value = converted
    if isinstance(converted, datetime.datetime):
        cell.number_format = "DD.MM.YYYY HH:MM"
    elif isinstance(converted, datetime.date):
        cell.number_format = "DD.MM.YYYY"


def cells_in(reference: str):
    min_col, min_row, max_col, max_row = range_boundaries(reference)
    for row in range(min_row, max_row + 1):
        for col in range(min_col, max_col + 1):
            yield f"{get_column_letter(col)}{row}"


def apply_style(sheet, reference: str, style: dict) -> None:
    for coordinate in cells_in(reference):
        cell = sheet[coordinate]
        font = copy(cell.font)
        for key, attribute in (("bold", "bold"), ("italic", "italic"), ("size", "size"), ("color", "color"), ("font", "name")):
            if key in style:
                setattr(font, attribute, style[key])
        cell.font = font
        if "fill" in style:
            cell.fill = PatternFill("solid", fgColor=style["fill"])
        if "align" in style or "wrap" in style or "valign" in style:
            cell.alignment = Alignment(
                horizontal=style.get("align", cell.alignment.horizontal),
                vertical=style.get("valign", cell.alignment.vertical),
                wrap_text=style.get("wrap", cell.alignment.wrap_text),
            )
        if style.get("border"):
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
        if "numberFormat" in style:
            cell.number_format = style["numberFormat"]


def auto_width(sheet) -> None:
    widths: dict[int, int] = {}
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            text = cell.value.strftime("%d.%m.%Y") if isinstance(cell.value, (datetime.date, datetime.datetime)) else str(cell.value)
            if text.startswith("="):
                text = "0" * 12
            longest = max(len(line) for line in text.split("\n"))
            widths[cell.column] = max(widths.get(cell.column, 0), longest)
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = min(max(width + 2, 8), 60)


def fill_sheet(sheet, spec: dict) -> None:
    rows = spec.get("rows") or []
    start_col, start_row = 1, 1
    if spec.get("start"):
        letters, start_row = coordinate_from_string(spec["start"])
        start_col = column_index_from_string(letters)
    parse_numbers = bool(spec.get("parseNumbers", False))
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            if value is None:
                continue
            put(sheet, f"{get_column_letter(start_col + c)}{start_row + r}", value, parse_numbers)
    for coordinate, value in (spec.get("cells") or {}).items():
        put(sheet, coordinate, value)

    width = max((len(row) for row in rows), default=0)
    if rows and spec.get("header", True) and width:
        header = f"{get_column_letter(start_col)}{start_row}:{get_column_letter(start_col + width - 1)}{start_row}"
        for coordinate in cells_in(header):
            font = copy(sheet[coordinate].font)
            font.bold = True
            sheet[coordinate].font = font
            sheet[coordinate].fill = HEADER_FILL
            sheet[coordinate].border = Border(bottom=THIN)
        if "freeze" not in spec:
            sheet.freeze_panes = f"{get_column_letter(start_col)}{start_row + 1}"
        if spec.get("autofilter", True) is True and len(rows) > 1:
            sheet.auto_filter.ref = f"{get_column_letter(start_col)}{start_row}:{get_column_letter(start_col + width - 1)}{start_row + len(rows) - 1}"
    if isinstance(spec.get("autofilter"), str):
        sheet.auto_filter.ref = spec["autofilter"]
    if "freeze" in spec:
        sheet.freeze_panes = spec["freeze"] or None

    for reference, number_format in (spec.get("numberFormats") or {}).items():
        for coordinate in cells_in(reference):
            sheet[coordinate].number_format = number_format
    for reference, style in (spec.get("styles") or {}).items():
        apply_style(sheet, reference, style)
    for reference in spec.get("merge") or []:
        sheet.merge_cells(reference)

    widths = spec.get("columnWidths", "auto")
    if widths == "auto":
        auto_width(sheet)
    elif isinstance(widths, dict):
        for column, value in widths.items():
            sheet.column_dimensions[column].width = value
    for row, height in (spec.get("rowHeights") or {}).items():
        sheet.row_dimensions[int(row)].height = height


def from_csv(path: str, delimiter: str | None) -> dict:
    text = Path(path).read_text(encoding="utf-8-sig", errors="replace")
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    rows = [row for row in csv.reader(io.StringIO(text), delimiter=delimiter)]
    return {"sheets": [{"name": Path(path).stem[:31] or "Sheet1", "rows": rows, "parseNumbers": True}]}


def main() -> int:
    parser = argparse.ArgumentParser(description="Write a workbook from a JSON spec or a CSV file.")
    parser.add_argument("spec", nargs="?", help="a JSON spec file, or - for stdin")
    parser.add_argument("output")
    parser.add_argument("--base", help="an existing workbook to change instead of starting empty")
    parser.add_argument("--csv", help="build the workbook from this CSV file")
    parser.add_argument("--delimiter", default=None)
    args = parser.parse_args()

    if args.csv:
        spec = from_csv(args.csv, args.delimiter)
    elif args.spec:
        text = sys.stdin.read() if args.spec == "-" else Path(args.spec).read_text(encoding="utf-8-sig")
        spec = json.loads(text)
    else:
        print("Give a JSON spec (or -) or --csv.", file=sys.stderr)
        return 2

    output = Path(args.output)
    if args.base:
        book = load_workbook(args.base, keep_vba=args.base.lower().endswith(".xlsm"))
    else:
        book = Workbook()
        book.remove(book.active)

    for sheet_spec in spec.get("sheets") or []:
        name = str(sheet_spec.get("name") or f"Sheet{len(book.worksheets) + 1}")[:31]
        sheet = book[name] if name in book.sheetnames else book.create_sheet(name)
        fill_sheet(sheet, sheet_spec)
    for name in spec.get("deleteSheets") or []:
        if name in book.sheetnames and len(book.sheetnames) > 1:
            book.remove(book[name])
    for name, reference in (spec.get("names") or {}).items():
        book.defined_names[name] = DefinedName(name, attr_text=reference)
    if not book.worksheets:
        book.create_sheet("Sheet1")
    if spec.get("active") in book.sheetnames:
        book.active = book.sheetnames.index(spec["active"])

    output.parent.mkdir(parents=True, exist_ok=True)
    book.save(output)
    formulas = sum(1 for sheet in book.worksheets for row in sheet.iter_rows() for cell in row if cell.data_type == "f")
    print(json.dumps({"output": str(output.resolve()), "sheets": book.sheetnames, "formulas": formulas}, ensure_ascii=False))
    if formulas:
        print("Formulas are stored without cached values; run check_formulas.py to verify references.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
