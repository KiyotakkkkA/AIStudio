# /// script
# requires-python = ">=3.10"
# dependencies = ["openpyxl==3.1.5"]
# ///
import argparse
import json
import re
import sys

from openpyxl import load_workbook

ERRORS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!")
SHEET_REFERENCE = re.compile(r"(?:'((?:[^']|'')+)'|([A-Za-z_Ѐ-ӿ][\w.Ѐ-ӿ]*))!")
STRING = re.compile(r'"(?:[^"]|"")*"')
CELL = re.compile(r"(?<![A-Za-z_!$])\$?([A-Z]{1,3})\$?(\d+)(?![\d(])")
CYRILLIC_FUNCTION = re.compile(r"[Ѐ-ӿ]+\s*\(")


def problems_of(formula: str, sheet_names: set[str], here: str) -> list[str]:
    found = []
    body = STRING.sub('""', formula)
    if "#REF!" in body:
        found.append("contains #REF!: a referenced cell or sheet was deleted")
    if body.count("(") != body.count(")"):
        found.append("unbalanced parentheses")
    if ";" in body:
        found.append("uses ';' as a separator; openpyxl formulas need ','")
    if CYRILLIC_FUNCTION.search(body):
        found.append("localized function name; use the English name, e.g. SUM instead of СУММ")
    for quoted, plain in SHEET_REFERENCE.findall(body):
        name = quoted.replace("''", "'") if quoted else plain
        if name not in sheet_names:
            found.append(f"refers to a missing sheet '{name}'")
    stripped = SHEET_REFERENCE.sub("", body)
    if not any(char == "!" for char in body):
        for column, row in CELL.findall(stripped):
            if f"{column}{row}" == here:
                found.append("refers to its own cell (circular reference)")
                break
    return found


def main() -> int:
    parser = argparse.ArgumentParser(description="Check every formula of a workbook.")
    parser.add_argument("file")
    parser.add_argument("--limit", type=int, default=100, help="how many problems to list")
    args = parser.parse_args()

    formulas_book = load_workbook(args.file, data_only=False)
    values_book = load_workbook(args.file, data_only=True)
    names = set(formulas_book.sheetnames)
    problems = []
    total = 0
    without_value = 0
    for sheet in formulas_book.worksheets:
        values = values_book[sheet.title]
        for row in sheet.iter_rows():
            for cell in row:
                value = values[cell.coordinate].value
                if cell.data_type == "f":
                    total += 1
                    formula = str(cell.value)
                    if value is None:
                        without_value += 1
                    for problem in problems_of(formula, names, cell.coordinate):
                        problems.append({"cell": f"{sheet.title}!{cell.coordinate}", "formula": formula, "problem": problem})
                if isinstance(value, str) and value in ERRORS:
                    problems.append({"cell": f"{sheet.title}!{cell.coordinate}", "formula": str(cell.value), "problem": f"calculated as {value}"})

    report = {
        "formulas": total,
        "withoutCachedValue": without_value,
        "problems": len(problems),
        "details": problems[: args.limit],
    }
    if total and without_value == total:
        report["note"] = "No formula has a cached value: errors that only appear on calculation are not visible until the file is opened in Excel."
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 1 if problems else 0


if __name__ == "__main__":
    sys.exit(main())
