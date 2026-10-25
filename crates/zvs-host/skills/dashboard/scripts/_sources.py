import csv
import datetime as dt
import re
import sqlite3
from pathlib import Path

TEXT_EXT = {".csv", ".tsv", ".txt", ".tab"}
JSON_EXT = {".json"}
JSONL_EXT = {".jsonl", ".ndjson"}
EXCEL_EXT = {".xlsx", ".xlsm"}
SQLITE_EXT = {".db", ".sqlite", ".sqlite3"}
PARQUET_EXT = {".parquet", ".pq"}
NAME = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")


class SourceError(Exception):
    pass


def ident(name: str) -> str:
    return '"' + name.replace('"', '""') + '"'


def literal(value) -> str:
    return "'" + str(value).replace("'", "''") + "'"


def split_source(text: str) -> tuple[str, str | None]:
    if "#" in text:
        path, part = text.rsplit("#", 1)
        if part and "/" not in part and "\\" not in part and Path(path).suffix:
            return path, part
    return text, None


def parse_named(text: str) -> tuple[str, str]:
    if "=" not in text:
        raise SourceError(f"expected NAME=PATH, got {text!r}")
    name, path = text.split("=", 1)
    name = name.strip()
    if not NAME.match(name):
        raise SourceError(f"source name {name!r} must be a plain identifier such as sales or orders_2026")
    return name, path.strip().strip('"')


def utf8_copy(path: Path, work: Path) -> Path:
    raw = path.read_bytes()
    try:
        raw.decode("utf-8")
        return path
    except UnicodeDecodeError:
        pass
    for encoding in ("cp1251", "utf-16", "latin-1"):
        try:
            text = raw.decode(encoding)
        except UnicodeDecodeError:
            continue
        target = work / f"{path.stem}.utf8{path.suffix}"
        target.write_text(text, encoding="utf-8")
        return target
    raise SourceError(f"{path}: unknown text encoding")


def sniff_delimiter(path: Path) -> str:
    if path.suffix.lower() in {".tsv", ".tab"}:
        return "\t"
    with path.open(encoding="utf-8-sig", errors="replace") as handle:
        sample = handle.read(65536)
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        return ","


def cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        if value.time() == dt.time(0, 0):
            return value.date().isoformat()
        return value.isoformat(sep="T", timespec="seconds")
    if isinstance(value, (dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def excel_rows(path: Path, sheet: str | None) -> tuple[list[str], list[list]]:
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=True)
    if sheet is None:
        target = book.worksheets[0]
    elif sheet in book.sheetnames:
        target = book[sheet]
    else:
        raise SourceError(f"{path}: no sheet {sheet!r}; sheets: {', '.join(book.sheetnames)}")
    rows = [list(row) for row in target.iter_rows(values_only=True)]
    book.close()
    while rows and all(value is None or str(value).strip() == "" for value in rows[0]):
        rows.pop(0)
    if not rows:
        return [], []
    header = [cell_text(value).strip() or f"column_{index + 1}" for index, value in enumerate(rows[0])]
    body = [row for row in rows[1:] if any(value is not None and str(value).strip() != "" for value in row)]
    return header, body


def excel_sheets(path: Path) -> list[str]:
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True)
    names = list(book.sheetnames)
    book.close()
    return names


def sqlite_tables(path: Path) -> list[str]:
    with sqlite3.connect(f"file:{path}?mode=ro", uri=True) as connection:
        rows = connection.execute(
            "SELECT name FROM sqlite_master WHERE type IN ('table', 'view') AND name NOT LIKE 'sqlite_%' ORDER BY name"
        ).fetchall()
    return [row[0] for row in rows]


def sqlite_rows(path: Path, table: str) -> tuple[list[str], list[list]]:
    with sqlite3.connect(f"file:{path}?mode=ro", uri=True) as connection:
        if table not in sqlite_tables(path):
            raise SourceError(f"{path}: no table {table!r}; tables: {', '.join(sqlite_tables(path))}")
        cursor = connection.execute(f"SELECT * FROM {ident(table)}")
        header = [column[0] for column in cursor.description]
        return header, [list(row) for row in cursor.fetchall()]


def write_csv(target: Path, header: list[str], rows: list[list]) -> Path:
    with target.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(header)
        for row in rows:
            writer.writerow([cell_text(value) for value in row])
    return target


def csv_options(delimiter: str, decimal_comma: bool, date_format: str | None) -> str:
    options = [f"delim={literal(delimiter)}", "header=true", "auto_detect=true", "sample_size=-1"]
    if decimal_comma:
        options.append("decimal_separator=','")
    if date_format:
        options.append(f"dateformat={literal(date_format)}")
    return ", ".join(options)


def register(connection, name: str, source: str, work: Path, decimal_comma=False, date_format=None) -> dict:
    path_text, part = split_source(source)
    path = Path(path_text)
    if not path.is_file():
        raise SourceError(f"file not found: {path_text}")
    suffix = path.suffix.lower()
    view = ident(name)
    info = {"name": name, "path": str(path), "part": part}
    if suffix in TEXT_EXT:
        text = utf8_copy(path, work)
        delimiter = sniff_delimiter(text)
        decimal = decimal_comma or delimiter == ";"
        info["delimiter"] = delimiter
        connection.execute(
            f"CREATE OR REPLACE VIEW {view} AS SELECT * FROM read_csv({literal(text)}, "
            f"{csv_options(delimiter, decimal, date_format)})"
        )
    elif suffix in JSON_EXT:
        connection.execute(f"CREATE OR REPLACE VIEW {view} AS SELECT * FROM read_json_auto({literal(path)})")
    elif suffix in JSONL_EXT:
        connection.execute(
            f"CREATE OR REPLACE VIEW {view} AS SELECT * FROM read_json_auto({literal(path)}, "
            "format='newline_delimited')"
        )
    elif suffix in PARQUET_EXT:
        connection.execute(f"CREATE OR REPLACE VIEW {view} AS SELECT * FROM read_parquet({literal(path)})")
    elif suffix in EXCEL_EXT:
        header, rows = excel_rows(path, part)
        info["sheet"] = part or excel_sheets(path)[0]
        target = write_csv(work / f"{name}.csv", header, rows)
        connection.execute(
            f"CREATE OR REPLACE VIEW {view} AS SELECT * FROM read_csv({literal(target)}, "
            f"{csv_options(',', False, date_format)})"
        )
    elif suffix in SQLITE_EXT:
        if part is None:
            raise SourceError(
                f"{path}: name the table as {path}#TABLE; tables: {', '.join(sqlite_tables(path))}"
            )
        header, rows = sqlite_rows(path, part)
        target = write_csv(work / f"{name}.csv", header, rows)
        connection.execute(
            f"CREATE OR REPLACE VIEW {view} AS SELECT * FROM read_csv({literal(target)}, "
            f"{csv_options(',', False, date_format)})"
        )
    else:
        raise SourceError(
            f"{path}: unsupported file type {suffix or '(none)'}; supported: csv, tsv, txt, json, jsonl, "
            "parquet, xlsx, xlsm, sqlite"
        )
    return info


def json_value(value):
    if value is None or isinstance(value, (bool, int, str)):
        return value
    if isinstance(value, float):
        return value if value == value and value not in (float("inf"), float("-inf")) else None
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return cell_text(value)
    if isinstance(value, (list, tuple)):
        return [json_value(item) for item in value]
    if isinstance(value, dict):
        return {str(key): json_value(item) for key, item in value.items()}
    try:
        return float(value)
    except (TypeError, ValueError):
        return str(value)
