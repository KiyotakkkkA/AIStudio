# /// script
# requires-python = ">=3.11"
# dependencies = ["duckdb==2.0.0", "openpyxl==3.1.5"]
# ///
import argparse
import csv
import datetime as dt
import difflib
import html
import json
import os
import re
import sys
from pathlib import Path

DASHBOARDS_ENV = "ZVS_DASHBOARDS_DIR"
NAME_MAX = 120
UNSAFE_NAME = re.compile(r'[\x00-\x1f<>:"/\\|?*]+')
ROWS_DEFAULT = 50_000
ROWS_CEILING = 200_000
WARN_BYTES = 5 * 1024 * 1024
MAX_BYTES = 15 * 1024 * 1024
AGGREGATIONS = {"sum", "avg", "count", "min", "max", "median", "distinct", "first", "last"}
FORMATS = {
    "auto", "integer", "decimal", "percent", "percentValue", "compact",
    "date", "month", "year", "datetime", "text",
}
TYPES = {"number", "date", "string"}
COMMON = {"type", "title", "subtitle", "span", "note", "dataset", "height"}
KEYS = {
    "kpi": {"label", "value", "agg", "format", "compare", "compareLabel", "goodWhen", "deltaMode", "spark", "hero"},
    "meter": {"label", "value", "target", "agg", "format"},
    "line": {"x", "y", "labels", "series", "seriesOrder", "agg", "format", "xFormat", "xLabel", "sort", "endLabels"},
    "area": {"x", "y", "labels", "series", "seriesOrder", "agg", "format", "xFormat", "xLabel", "sort", "stacked", "endLabels"},
    "bar": {
        "x", "y", "labels", "series", "seriesOrder", "agg", "format", "xFormat", "xLabel",
        "sort", "limit", "stacked", "orientation", "valueLabels",
    },
    "donut": {"label", "value", "agg", "format", "limit", "order", "centerLabel", "labelTitle", "valueLabel"},
    "scatter": {"x", "y", "series", "label", "format", "xFormat", "xLabel", "yLabel"},
    "heatmap": {"x", "y", "value", "agg", "format", "xFormat", "yFormat", "xOrder", "yOrder", "cellHeight", "cellLabels", "sortCategories"},
    "funnel": {"label", "value", "agg", "format", "order", "labelTitle", "valueLabel"},
    "table": {"columns", "sort", "limit", "search"},
    "text": {"text"},
}
NUMBER = re.compile(r"^[+-]?(\d+([.,]\d*)?|[.,]\d+)([eE][+-]?\d+)?$")
ISO_DATE = re.compile(r"^\d{4}-\d{2}(-\d{2})?([ T]\d{2}:\d{2}(:\d{2}(\.\d+)?)?)?(Z|[+-]\d{2}:?\d{2})?$")
RU_DATE = re.compile(r"^(\d{2})\.(\d{2})\.(\d{4})(?:[ T](\d{2}):(\d{2})(?::(\d{2}))?)?$")
SPACES = re.compile(r"[\s  ]")


class Problems:
    def __init__(self) -> None:
        self.errors: list[str] = []
        self.warnings: list[str] = []

    def error(self, text: str) -> None:
        self.errors.append(text)

    def warn(self, text: str) -> None:
        self.warnings.append(text)


def file_name(name: str | None, title: str) -> str:
    base = (name or "").strip()
    if base.lower().endswith(".html"):
        base = base[:-5]
    if not base:
        base = f"{title} {dt.date.today().isoformat()}"
    base = UNSAFE_NAME.sub(" ", base)
    base = re.sub(r"\s+", " ", base).strip(" .")[:NAME_MAX].strip(" .")
    return f"{base or 'dashboard'}.html"


def assets_dir() -> Path:
    skill = os.environ.get("ZVS_SKILL_DIR")
    if skill:
        return Path(skill) / "assets"
    return Path(__file__).resolve().parent.parent / "assets"


def read_text(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1")


def load_csv(path: Path, delimiter: str | None) -> tuple[list[str], list[list], str]:
    text = read_text(path)
    if delimiter is None:
        if path.suffix.lower() in {".tsv", ".tab"}:
            delimiter = "\t"
        else:
            try:
                delimiter = csv.Sniffer().sniff(text[:65536], delimiters=",;\t|").delimiter
            except csv.Error:
                delimiter = ","
    rows = list(csv.reader(text.splitlines(), delimiter=delimiter))
    rows = [row for row in rows if any(cell.strip() for cell in row)]
    if not rows:
        return [], [], delimiter
    header = [cell.strip() or f"column_{index + 1}" for index, cell in enumerate(rows[0])]
    return header, rows[1:], delimiter


def from_records(records: list) -> tuple[list[str], list[list]]:
    header: list[str] = []
    for record in records:
        if not isinstance(record, dict):
            raise ValueError("expected a list of objects")
        for key in record:
            if key not in header:
                header.append(key)
    return header, [[record.get(key) for key in header] for record in records]


def load_json(path: Path) -> tuple[list[str], list[list]]:
    text = read_text(path)
    if path.suffix.lower() in {".jsonl", ".ndjson"}:
        return from_records([json.loads(line) for line in text.splitlines() if line.strip()])
    data = json.loads(text)
    if isinstance(data, dict) and isinstance(data.get("columns"), list) and isinstance(data.get("rows"), list):
        names = [column["name"] if isinstance(column, dict) else str(column) for column in data["columns"]]
        if data["rows"] and isinstance(data["rows"][0], dict):
            return from_records(data["rows"])
        return names, [list(row) for row in data["rows"]]
    if isinstance(data, dict):
        for key in ("rows", "data", "items", "records", "results"):
            if isinstance(data.get(key), list):
                return from_records(data[key])
        raise ValueError("expected a list of objects or an object with a rows list")
    return from_records(data)


def load_excel(path: Path, sheet: str | None) -> tuple[list[str], list[list]]:
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=True)
    target = book[sheet] if sheet else book.worksheets[0]
    rows = [list(row) for row in target.iter_rows(values_only=True)]
    book.close()
    rows = [row for row in rows if any(value is not None and str(value).strip() != "" for value in row)]
    if not rows:
        return [], []
    header = [str(value).strip() if value is not None else f"column_{index + 1}" for index, value in enumerate(rows[0])]
    return header, rows[1:]


def as_number(value, decimal_comma: bool):
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return value
    text = SPACES.sub("", str(value))
    if not NUMBER.match(text):
        return None
    if "," in text:
        if not decimal_comma or "." in text:
            return None
        text = text.replace(",", ".")
    try:
        number = float(text)
    except ValueError:
        return None
    return number


def as_date(value):
    if isinstance(value, dt.datetime):
        if value.time() == dt.time(0, 0):
            return value.date().isoformat()
        return value.replace(tzinfo=None).isoformat(timespec="seconds")
    if isinstance(value, dt.date):
        return value.isoformat()
    text = str(value).strip()
    if ISO_DATE.match(text):
        return text.replace(" ", "T")
    match = RU_DATE.match(text)
    if match:
        day, month, year, hour, minute, second = match.groups()
        date = f"{year}-{month}-{day}"
        if hour:
            return f"{date}T{hour}:{minute}:{second or '00'}"
        return date
    return None


def blank(value) -> bool:
    return value is None or (isinstance(value, str) and value.strip() == "")


def column_type(values: list, decimal_comma: bool) -> str:
    present = [value for value in values if not blank(value)]
    if not present:
        return "string"
    if all(as_number(value, decimal_comma) is not None for value in present):
        return "number"
    if all(as_date(value) is not None for value in present):
        return "date"
    return "string"


def convert(value, kind: str, decimal_comma: bool):
    if blank(value):
        return None
    if kind == "number":
        number = as_number(value, decimal_comma)
        if number is None:
            return None
        if isinstance(number, float) and number.is_integer() and abs(number) < 2**53:
            return int(number)
        return round(number, 10) if isinstance(number, float) else number
    if kind == "date":
        return as_date(value) or str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (dt.date, dt.datetime)):
        return as_date(value)
    return str(value).strip()


def load_sql(name: str, spec: dict) -> tuple[list[str], list[list]]:
    import tempfile

    import duckdb

    from _sources import SourceError, parse_named, register

    sources = spec.get("sources") or {}
    if not isinstance(sources, dict) or not sources:
        raise ValueError('give "sources": {"NAME": "ABSOLUTE_PATH"} for the tables the SQL reads')
    connection = duckdb.connect()
    try:
        with tempfile.TemporaryDirectory() as work:
            for source_name, path in sources.items():
                source_name, path = parse_named(f"{source_name}={path}")
                register(
                    connection,
                    source_name,
                    path,
                    Path(work),
                    spec.get("decimal") == ",",
                    spec.get("dateFormat"),
                )
            cursor = connection.execute(str(spec["sql"]).strip().rstrip(";"))
            header = [column[0] for column in cursor.description]
            return header, [list(row) for row in cursor.fetchall()]
    except (SourceError, duckdb.Error) as error:
        raise ValueError(f"{error}".strip()) from error


def load_dataset(name: str, spec, base: Path, problems: Problems) -> dict | None:
    if isinstance(spec, list):
        spec = {"rows": spec}
    if not isinstance(spec, dict):
        problems.error(f"datasets.{name}: expected an object with sql, path or rows")
        return None
    delimiter = None
    try:
        if "rows" in spec:
            header, rows = from_records(spec["rows"])
        elif "sql" in spec:
            header, rows = load_sql(name, spec)
        elif "path" in spec:
            path = Path(spec["path"])
            if not path.is_absolute():
                path = base / path
            if not path.is_file():
                problems.error(f"datasets.{name}: file not found: {path}")
                return None
            suffix = path.suffix.lower()
            if suffix in {".csv", ".tsv", ".txt", ".tab"}:
                header, rows, delimiter = load_csv(path, spec.get("delimiter"))
            elif suffix in {".json", ".jsonl", ".ndjson"}:
                header, rows = load_json(path)
            elif suffix in {".xlsx", ".xlsm"}:
                header, rows = load_excel(path, spec.get("sheet"))
            else:
                problems.error(
                    f"datasets.{name}: {suffix} is not supported here; convert it with query_data.py --out result.csv"
                )
                return None
        else:
            problems.error(f"datasets.{name}: give sql with sources, path or rows")
            return None
    except (ValueError, KeyError, OSError, json.JSONDecodeError) as error:
        problems.error(f"datasets.{name}: {error}")
        return None
    rename = spec.get("rename") or {}
    header = [rename.get(column, column) for column in header]
    if len(set(header)) != len(header):
        problems.error(f"datasets.{name}: duplicate column names {sorted({c for c in header if header.count(c) > 1})}")
        return None
    keep = spec.get("columns")
    indexes = list(range(len(header)))
    if keep:
        missing = [column for column in keep if column not in header]
        for column in missing:
            problems.error(f"datasets.{name}.columns: no column {column!r}{suggest(column, header)}")
        indexes = [header.index(column) for column in keep if column in header]
    limit = min(int(spec.get("maxRows", ROWS_DEFAULT)), ROWS_CEILING)
    if len(rows) > limit:
        problems.error(
            f"datasets.{name}: {len(rows)} rows, the limit is {limit}; aggregate with query_data.py "
            "to the level the dashboard and its filters need"
        )
        return None
    decimal_comma = spec.get("decimal") == "," or (spec.get("decimal") is None and delimiter == ";")
    width = len(header)
    rows = [list(row) + [None] * (width - len(row)) for row in rows]
    columns = []
    overrides = spec.get("types") or {}
    for index in indexes:
        column = header[index]
        kind = overrides.get(column) or column_type([row[index] for row in rows], decimal_comma)
        if kind not in TYPES:
            problems.error(f"datasets.{name}.types.{column}: use number, date or string")
            kind = "string"
        columns.append({"name": column, "type": kind, "index": index})
    data_rows = [[convert(row[column["index"]], column["type"], decimal_comma) for column in columns] for row in rows]
    if not data_rows:
        problems.warn(f"datasets.{name}: no rows")
    return {"columns": [{"name": c["name"], "type": c["type"]} for c in columns], "rows": data_rows}


def suggest(name: str, options) -> str:
    close = difflib.get_close_matches(str(name), list(options), n=1)
    return f" (did you mean {close[0]!r}?)" if close else ""


def check_field(path: str, field, dataset: str | None, datasets: dict, problems: Problems, numeric=False) -> None:
    if field is None:
        return
    if not isinstance(field, str):
        problems.error(f"{path}: expected a column name")
        return
    if dataset is None or dataset not in datasets:
        return
    types = {column["name"]: column["type"] for column in datasets[dataset]["columns"]}
    if field not in types:
        problems.error(f"{path}: dataset {dataset!r} has no column {field!r}{suggest(field, types)}")
    elif numeric and types[field] != "number":
        problems.warn(f"{path}: column {field!r} is {types[field]}, not a number; only count and distinct work on it")


def check_measure(path: str, spec, dataset, datasets, problems: Problems, required=True) -> None:
    if spec is None:
        if required:
            problems.error(f"{path}: required")
        return
    if isinstance(spec, (int, float)) and not isinstance(spec, bool):
        return
    if isinstance(spec, str):
        if dataset is None:
            problems.error(f"{path}: a column name needs the card's dataset")
            return
        check_field(path, spec, dataset, datasets, problems, numeric=True)
        return
    if not isinstance(spec, dict):
        problems.error(f"{path}: expected a number, a column name or an object with field and agg")
        return
    if "ratio" in spec:
        parts = spec["ratio"]
        if not isinstance(parts, list) or len(parts) != 2:
            problems.error(f"{path}.ratio: a list of two measures, numerator and denominator")
            return
        for index, part in enumerate(parts):
            check_measure(f"{path}.ratio[{index}]", part, spec.get("dataset", dataset), datasets, problems)
        return
    source = spec.get("dataset", dataset)
    if spec.get("dataset") is not None and spec["dataset"] not in datasets:
        problems.error(f"{path}.dataset: no dataset {spec['dataset']!r}{suggest(spec['dataset'], datasets)}")
        return
    if source is None:
        problems.error(f"{path}: name a dataset for this measure")
        return
    agg = spec.get("agg")
    if agg is not None and agg not in AGGREGATIONS:
        problems.error(f"{path}.agg: {agg!r} is not one of {', '.join(sorted(AGGREGATIONS))}")
    if spec.get("field") is None and agg not in (None, "count"):
        problems.error(f"{path}: field is required unless agg is count")
    check_field(f"{path}.field", spec.get("field"), source, datasets, problems, numeric=agg not in ("count", "distinct", "first", "last"))
    check_field(f"{path}.orderBy", spec.get("orderBy"), source, datasets, problems)
    if spec.get("period") is not None:
        if spec["period"] not in ("last", "previous"):
            problems.error(f"{path}.period: last or previous")
        if spec.get("by") is None:
            problems.error(f"{path}.by: name the period column, for example month")
    check_field(f"{path}.by", spec.get("by"), source, datasets, problems)
    for key in (spec.get("where") or {}):
        check_field(f"{path}.where.{key}", key, source, datasets, problems)


def check_format(path: str, value, problems: Problems) -> None:
    if value is None:
        return
    if not isinstance(value, str) or not (value in FORMATS or re.fullmatch(r"currency:[A-Z]{3}", value)):
        problems.error(f"{path}: unknown format {value!r}; use {', '.join(sorted(FORMATS))} or currency:RUB")


def check_component(path: str, item, columns: int, datasets: dict, used: set, problems: Problems) -> None:
    if not isinstance(item, dict):
        problems.error(f"{path}: expected an object")
        return
    kind = item.get("type")
    if kind not in KEYS:
        problems.error(f"{path}.type: {kind!r} is not one of {', '.join(sorted(KEYS))}")
        return
    for key in item:
        if key not in COMMON and key not in KEYS[kind]:
            problems.warn(f"{path}.{key}: not used by {kind}{suggest(key, KEYS[kind] | COMMON)}")
    dataset = item.get("dataset")
    needs_data = kind not in ("text",) and not (
        kind in ("kpi", "meter") and all(isinstance(item.get(key), (int, float)) for key in ("value", "target") if key in item)
    )
    if dataset is not None:
        if dataset not in datasets:
            problems.error(f"{path}.dataset: no dataset {dataset!r}{suggest(dataset, datasets)}")
            return
        used.add(dataset)
    elif needs_data and kind not in ("kpi", "meter"):
        problems.error(f"{path}.dataset: required")
        return
    span = item.get("span")
    if span is not None and (not isinstance(span, int) or not 1 <= span <= columns):
        problems.error(f"{path}.span: must be an integer from 1 to {columns}")
    height = item.get("height")
    if height is not None and (not isinstance(height, (int, float)) or not 120 <= height <= 900):
        problems.error(f"{path}.height: must be between 120 and 900")
    for key in ("format", "xFormat", "yFormat"):
        check_format(f"{path}.{key}", item.get(key), problems)
    agg = item.get("agg")
    if agg is not None and agg not in AGGREGATIONS:
        problems.error(f"{path}.agg: {agg!r} is not one of {', '.join(sorted(AGGREGATIONS))}")

    if kind in ("line", "area", "bar"):
        check_field(f"{path}.x", item.get("x"), dataset, datasets, problems)
        if item.get("x") is None:
            problems.error(f"{path}.x: required")
        ys = item.get("y")
        ys = ys if isinstance(ys, list) else ([] if ys is None else [ys])
        if not ys and agg != "count":
            problems.error(f"{path}.y: required unless agg is count")
        if item.get("series") and len(ys) > 1:
            problems.error(f"{path}: use either series or several y columns, not both")
        if len(ys) > 8:
            problems.error(f"{path}.y: at most 8 series; fold the rest or split the chart")
        for index, y in enumerate(ys):
            field = y.get("field") if isinstance(y, dict) else y
            check_field(f"{path}.y[{index}]", field, dataset, datasets, problems, numeric=agg not in ("count", "distinct"))
        check_field(f"{path}.series", item.get("series"), dataset, datasets, problems)
        if kind == "bar" and item.get("orientation") not in (None, "vertical", "horizontal"):
            problems.error(f"{path}.orientation: vertical or horizontal")
        if item.get("sort") not in (None, "desc", "asc", "x", "none"):
            problems.error(f"{path}.sort: desc, asc, x or none")
    elif kind == "donut":
        check_field(f"{path}.label", item.get("label"), dataset, datasets, problems)
        if item.get("label") is None:
            problems.error(f"{path}.label: required")
        check_field(f"{path}.value", item.get("value"), dataset, datasets, problems, numeric=True)
        if item.get("value") is None and agg != "count":
            problems.error(f"{path}.value: required unless agg is count")
        if isinstance(item.get("limit"), int) and item["limit"] > 6:
            problems.warn(f"{path}.limit: a donut shows at most 6 segments; the rest folds into one")
    elif kind == "scatter":
        for key in ("x", "y"):
            if item.get(key) is None:
                problems.error(f"{path}.{key}: required")
            check_field(f"{path}.{key}", item.get(key), dataset, datasets, problems, numeric=True)
        check_field(f"{path}.series", item.get("series"), dataset, datasets, problems)
        check_field(f"{path}.label", item.get("label"), dataset, datasets, problems)
    elif kind == "heatmap":
        for key in ("x", "y"):
            if item.get(key) is None:
                problems.error(f"{path}.{key}: required")
            check_field(f"{path}.{key}", item.get(key), dataset, datasets, problems)
        check_field(f"{path}.value", item.get("value"), dataset, datasets, problems, numeric=agg not in ("count", "distinct"))
        if item.get("value") is None and agg != "count":
            problems.error(f"{path}.value: required unless agg is count")
    elif kind == "funnel":
        check_field(f"{path}.label", item.get("label"), dataset, datasets, problems)
        check_field(f"{path}.value", item.get("value"), dataset, datasets, problems, numeric=True)
        if item.get("label") is None:
            problems.error(f"{path}.label: required")
    elif kind == "kpi":
        check_measure(f"{path}.value", item.get("value"), dataset, datasets, problems)
        check_measure(f"{path}.compare", item.get("compare"), dataset, datasets, problems, required=False)
        if item.get("goodWhen") not in (None, "up", "down"):
            problems.error(f"{path}.goodWhen: up or down")
        if item.get("deltaMode") not in (None, "percent", "absolute"):
            problems.error(f"{path}.deltaMode: percent or absolute")
        spark = item.get("spark")
        if spark is not None:
            if not isinstance(spark, dict) or dataset is None:
                problems.error(f"{path}.spark: an object with x and y on the card's dataset")
            else:
                check_field(f"{path}.spark.x", spark.get("x"), dataset, datasets, problems)
                check_field(f"{path}.spark.y", spark.get("y"), dataset, datasets, problems, numeric=True)
        if not item.get("label") and not item.get("title"):
            problems.error(f"{path}.label: a stat tile needs a label")
    elif kind == "meter":
        check_measure(f"{path}.value", item.get("value"), dataset, datasets, problems)
        check_measure(f"{path}.target", item.get("target"), dataset, datasets, problems)
    elif kind == "table":
        for index, column in enumerate(item.get("columns") or []):
            field = column.get("field") if isinstance(column, dict) else column
            check_field(f"{path}.columns[{index}]", field, dataset, datasets, problems)
            if isinstance(column, dict):
                check_format(f"{path}.columns[{index}].format", column.get("format"), problems)
        sort = item.get("sort")
        if sort is not None:
            if not isinstance(sort, dict) or sort.get("dir") not in ("asc", "desc"):
                problems.error(f"{path}.sort: {{\"field\": ..., \"dir\": \"asc\" | \"desc\"}}")
            else:
                check_field(f"{path}.sort.field", sort.get("field"), dataset, datasets, problems)
    elif kind == "text":
        if not isinstance(item.get("text"), str) or not item["text"].strip():
            problems.error(f"{path}.text: required")


def validate(spec: dict, datasets: dict, problems: Problems) -> None:
    if not isinstance(spec.get("title"), str) or not spec["title"].strip():
        problems.error("title: required")
    if spec.get("lang", "ru") not in ("ru", "en"):
        problems.error("lang: ru or en")
    if spec.get("theme", "auto") not in ("auto", "light", "dark"):
        problems.error("theme: auto, light or dark")
    sections = spec.get("sections")
    if not isinstance(sections, list) or not sections:
        problems.error("sections: at least one section with items")
        return
    used: set = set()
    count = 0
    for index, section in enumerate(sections):
        path = f"sections[{index}]"
        if not isinstance(section, dict) or not isinstance(section.get("items"), list) or not section["items"]:
            problems.error(f"{path}.items: at least one component")
            continue
        columns = section.get("columns", 4)
        if not isinstance(columns, int) or not 1 <= columns <= 4:
            problems.error(f"{path}.columns: an integer from 1 to 4")
            columns = 4
        for position, item in enumerate(section["items"]):
            check_component(f"{path}.items[{position}]", item, columns, datasets, used, problems)
            count += 1
    for index, item in enumerate(spec.get("filters") or []):
        path = f"filters[{index}]"
        if not isinstance(item, dict) or item.get("type", "select") not in ("select", "dateRange"):
            problems.error(f"{path}.type: select or dateRange")
            continue
        field = item.get("field")
        owners = [name for name, data in datasets.items() if any(c["name"] == field for c in data["columns"])]
        if not owners:
            names = {c["name"] for data in datasets.values() for c in data["columns"]}
            problems.error(f"{path}.field: no dataset has a column {field!r}{suggest(field, names)}")
            continue
        if item.get("type") == "dateRange":
            kinds = {c["type"] for name in owners for c in datasets[name]["columns"] if c["name"] == field}
            if "date" not in kinds:
                problems.error(f"{path}: dateRange needs a date column; {field!r} is {', '.join(sorted(kinds))}")
        for name in item.get("datasets") or []:
            if name not in datasets:
                problems.error(f"{path}.datasets: no dataset {name!r}")
    if count > 40:
        problems.warn(f"{count} components; a dashboard over about 30 blocks is hard to read, consider splitting it")
    for name in datasets:
        if name not in used and not any(name in (f.get("datasets") or []) for f in spec.get("filters") or []):
            problems.warn(f"datasets.{name}: not used by any component")


def payload(spec: dict, datasets: dict) -> dict:
    result = {key: value for key, value in spec.items() if key != "datasets"}
    result["lang"] = spec.get("lang", "ru")
    result["theme"] = spec.get("theme", "auto")
    result.setdefault("generatedAt", dt.datetime.now().replace(microsecond=0).isoformat(timespec="minutes"))
    result["datasets"] = datasets
    return result


def script_json(data: dict) -> str:
    text = json.dumps(data, ensure_ascii=False, separators=(",", ":"))
    return text.replace("</", "<\\/").replace("<!--", "<\\!--")


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Validate a dashboard spec, load its datasets and write one self-contained HTML file."
    )
    parser.add_argument(
        "spec",
        nargs="?",
        default="-",
        help="absolute path of the dashboard JSON spec, or - to read it from standard input",
    )
    parser.add_argument(
        "--name",
        help="file name of the dashboard without folders, for example «Продажи 2026»; the title and today's date by default",
    )
    parser.add_argument(
        "--out",
        help="absolute path of an extra copy; outside the app it is the only output",
    )
    args = parser.parse_args()
    library = os.environ.get(DASHBOARDS_ENV)
    if not library and not args.out:
        print("--out is required when the dashboards folder is not set", file=sys.stderr)
        return 2

    from_stdin = args.spec == "-"
    spec_path = Path.cwd() / "spec.json" if from_stdin else Path(args.spec)
    try:
        if from_stdin:
            text = sys.stdin.buffer.read().decode("utf-8-sig")
            if not text.strip():
                raise ValueError("nothing on standard input; pass the spec JSON as stdin or give a file path")
        else:
            text = read_text(spec_path)
        spec = json.loads(text)
        if not isinstance(spec, dict):
            raise ValueError("the spec must be a JSON object")
    except (OSError, ValueError) as error:
        print(json.dumps({"ok": False, "problems": [f"spec: {error}"]}, ensure_ascii=False))
        return 1
    problems = Problems()
    datasets = {}
    for name, dataset in (spec.get("datasets") or {}).items():
        loaded = load_dataset(name, dataset, spec_path.parent, problems)
        if loaded is not None:
            datasets[name] = loaded
    if not spec.get("datasets"):
        problems.warn("datasets: none; only text and literal numbers can be shown")
    validate(spec, datasets, problems)
    if problems.errors:
        print(json.dumps({"ok": False, "problems": problems.errors, "warnings": problems.warnings}, ensure_ascii=False, indent=1))
        return 1

    assets = assets_dir()
    template = (assets / "template.html").read_text(encoding="utf-8")
    page = (
        template.replace("{{lang}}", html.escape(spec.get("lang", "ru")))
        .replace("{{title}}", html.escape(spec["title"]))
        .replace("{{css}}", (assets / "zvs-dash.css").read_text(encoding="utf-8"))
        .replace("{{js}}", (assets / "zvs-dash.js").read_text(encoding="utf-8"))
        .replace("{{data}}", script_json(payload(spec, datasets)))
    )
    size = len(page.encode("utf-8"))
    if size > MAX_BYTES:
        print(json.dumps({"ok": False, "problems": [f"the page is {size // 1024} KiB; aggregate the data further"]}, ensure_ascii=False))
        return 1
    if size > WARN_BYTES:
        problems.warn(f"the page is {size // 1024} KiB; it opens slowly, consider aggregating")
    targets = []
    if library:
        targets.append(Path(library) / file_name(args.name, spec["title"]))
    if args.out:
        targets.append(Path(args.out))
    for target in targets:
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(page, encoding="utf-8")
    report = {
        "ok": True,
        "out": str(targets[0]),
        "kib": round(size / 1024, 1),
        "datasets": {
            name: {"rows": len(data["rows"]), "columns": {c["name"]: c["type"] for c in data["columns"]}}
            for name, data in datasets.items()
        },
        "components": sum(len(section.get("items", [])) for section in spec["sections"]),
        "warnings": problems.warnings,
    }
    if len(targets) > 1:
        report["copy"] = str(targets[1])
    print(json.dumps(report, ensure_ascii=False, indent=1))
    return 0


if __name__ == "__main__":
    sys.exit(main())
